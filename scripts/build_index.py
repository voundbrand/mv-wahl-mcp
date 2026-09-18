#!/usr/bin/env python3
"""Build SQLite FTS5 index from programme texts + Wahl-O-Mat dataset."""

from __future__ import annotations

import json
import re
import sqlite3
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from mv_wahl_mcp.paths import (  # noqa: E402
    INDEX_DB,
    PARTIES_JSON,
    PROGRAMMES_TEXT,
    WAHLOMAT_DIR,
)

CHUNK_SIZE = 1200
CHUNK_OVERLAP = 200


def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]

    chunks: list[str] = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + size, n)
        if end < n:
            # Prefer break at paragraph / sentence / space
            window = text[start:end]
            for sep in ("\n\n", "\n", ". ", "; ", ", ", " "):
                pos = window.rfind(sep)
                if pos > size // 3:
                    end = start + pos + len(sep)
                    break
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= n:
            break
        start = max(end - overlap, start + 1)
    return chunks


def create_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        DROP TABLE IF EXISTS chunks_fts;
        DROP TABLE IF EXISTS chunks;
        DROP TABLE IF EXISTS parties;
        DROP TABLE IF EXISTS wahlomat_rows;

        CREATE TABLE parties (
            id TEXT PRIMARY KEY,
            short_name TEXT NOT NULL,
            full_name TEXT NOT NULL,
            source_url TEXT,
            source_note TEXT,
            mirror_url TEXT,
            pdf_file TEXT,
            text_file TEXT,
            char_count INTEGER,
            chunk_count INTEGER,
            indexed INTEGER NOT NULL DEFAULT 0
        );

        CREATE TABLE chunks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            party_id TEXT NOT NULL,
            corpus TEXT NOT NULL,          -- 'programme' | 'wahlomat'
            chunk_index INTEGER NOT NULL,
            title TEXT,
            text TEXT NOT NULL,
            FOREIGN KEY(party_id) REFERENCES parties(id)
        );

        CREATE VIRTUAL TABLE chunks_fts USING fts5(
            party_id,
            corpus,
            title,
            text,
            content='chunks',
            content_rowid='id',
            tokenize='unicode61 remove_diacritics 2'
        );

        CREATE TRIGGER chunks_ai AFTER INSERT ON chunks BEGIN
            INSERT INTO chunks_fts(rowid, party_id, corpus, title, text)
            VALUES (new.id, new.party_id, new.corpus, new.title, new.text);
        END;
        CREATE TRIGGER chunks_ad AFTER DELETE ON chunks BEGIN
            INSERT INTO chunks_fts(chunks_fts, rowid, party_id, corpus, title, text)
            VALUES ('delete', old.id, old.party_id, old.corpus, old.title, old.text);
        END;
        CREATE TRIGGER chunks_au AFTER UPDATE ON chunks BEGIN
            INSERT INTO chunks_fts(chunks_fts, rowid, party_id, corpus, title, text)
            VALUES ('delete', old.id, old.party_id, old.corpus, old.title, old.text);
            INSERT INTO chunks_fts(rowid, party_id, corpus, title, text)
            VALUES (new.id, new.party_id, new.corpus, new.title, new.text);
        END;
        """
    )


def index_programmes(conn: sqlite3.Connection, parties: list[dict]) -> None:
    for p in parties:
        text_path = PROGRAMMES_TEXT / p["text_file"]
        text = text_path.read_text(encoding="utf-8", errors="replace") if text_path.exists() else ""
        chunks = chunk_text(text)
        conn.execute(
            """
            INSERT INTO parties (
                id, short_name, full_name, source_url, source_note, mirror_url,
                pdf_file, text_file, char_count, chunk_count, indexed
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                p["id"],
                p["short_name"],
                p["full_name"],
                p.get("source_url"),
                p.get("source_note"),
                p.get("mirror_url"),
                p.get("pdf_file"),
                p.get("text_file"),
                len(text),
                len(chunks),
                1 if chunks else 0,
            ),
        )
        for i, chunk in enumerate(chunks):
            conn.execute(
                """
                INSERT INTO chunks (party_id, corpus, chunk_index, title, text)
                VALUES (?, 'programme', ?, ?, ?)
                """,
                (p["id"], i, f"{p['short_name']} Wahlprogramm §{i + 1}", chunk),
            )
        print(f"  programme {p['short_name']}: {len(text)} chars -> {len(chunks)} chunks")


def index_wahlomat(conn: sqlite3.Connection) -> None:
    try:
        import openpyxl
    except ImportError:
        print("  openpyxl missing; skip Wahl-O-Mat")
        return

    xlsx = next(WAHLOMAT_DIR.glob("*Datensatz*.xlsx"), None)
    if not xlsx:
        print("  no Wahl-O-Mat xlsx found; skip")
        return

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb["Datensatz MV 2026"] if "Datensatz MV 2026" in wb.sheetnames else wb[wb.sheetnames[-1]]

    # Map short names / ids
    party_rows = conn.execute("SELECT id, short_name FROM parties").fetchall()
    by_short = {r[1].casefold(): r[0] for r in party_rows}
    # aliases
    aliases = {
        "grüne": "GRUENE",
        "gruene": "GRUENE",
        "die linke": "Die_Linke",
        "die partei": "Die_PARTEI",
        "freie wähler": "FREIE_WAEHLER",
        "ödp": "OEDP",
        "oedp": "OEDP",
        "bündnis c": "Buendnis_C",
        "buendnis c": "Buendnis_C",
        "handwerker partei deutschland": "Handwerker",
        "wir leben demokratie": "WLD",
        "wir leben demokratie (wld)": "WLD",
        "team freiheit": "Team_Freiheit",
        "piraten": "PIRATEN",
        "pdf": "PdF",
        "volt": "Volt",
        "bsw": "BSW",
        "kpd": "KPD",
        "spd": "SPD",
        "afd": "AfD",
        "cdu": "CDU",
        "fdp": "FDP",
        "tierschutzpartei": "Tierschutzpartei",
    }

    def resolve_party(short: str) -> str | None:
        if not short:
            return None
        key = short.casefold().strip()
        if key in by_short:
            return by_short[key]
        if key in aliases:
            return aliases[key]
        # fuzzy contains
        for s, pid in by_short.items():
            if key in s or s in key:
                return pid
        return None

    count = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if not row or row[1] is None:
            continue
        short = str(row[1]).strip()
        party_id = resolve_party(short)
        if not party_id:
            print(f"  WARN unmatched party: {short!r}")
            continue
        these_nr = row[3]
        these_title = row[4] or ""
        these_text = row[5] or ""
        position = row[6] or ""
        begruendung = row[7] or ""
        title = f"Wahl-O-Mat These {these_nr}: {these_title}"
        body = (
            f"These: {these_text}\n"
            f"Position der Partei: {position}\n"
            f"Begründung: {begruendung}"
        )
        conn.execute(
            """
            INSERT INTO chunks (party_id, corpus, chunk_index, title, text)
            VALUES (?, 'wahlomat', ?, ?, ?)
            """,
            (party_id, int(these_nr) if these_nr is not None else i, title, body),
        )
        count += 1
    print(f"  wahlomat: {count} position rows indexed")


def main() -> None:
    parties = json.loads(PARTIES_JSON.read_text(encoding="utf-8"))
    INDEX_DB.parent.mkdir(parents=True, exist_ok=True)
    if INDEX_DB.exists():
        INDEX_DB.unlink()

    conn = sqlite3.connect(INDEX_DB)
    try:
        create_schema(conn)
        print("Indexing programmes...")
        index_programmes(conn, parties)
        print("Indexing Wahl-O-Mat...")
        index_wahlomat(conn)
        conn.commit()
        n_parties = conn.execute("SELECT COUNT(*) FROM parties WHERE indexed=1").fetchone()[0]
        n_chunks = conn.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
        print(f"Done. {n_parties}/19 parties indexed, {n_chunks} chunks total -> {INDEX_DB}")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
